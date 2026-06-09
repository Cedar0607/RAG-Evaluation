from __future__ import annotations

import argparse
import ast
import asyncio
import importlib
import json
import math
import sys
import types
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook


REQUIRED_COLUMNS = ("question", "answer", "contexts", "ground_truths")
METRIC_COLUMNS = (
    "context_precision",
    "context_recall",
    "faithfulness",
    "factual_correctness",
    "ragas_error",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run a minimal Ragas evaluation from an Excel file.")
    parser.add_argument("--input", required=True, type=Path, help="Input Excel path.")
    parser.add_argument(
        "--config",
        type=Path,
        default=Path("ragas_config.json"),
        help="Judge LLM config path.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("output/ragas_result.xlsx"),
        help="Output Excel path.",
    )
    parser.add_argument("--sheet", default=None, help="Sheet name. Defaults to the active sheet.")
    parser.add_argument("--limit", type=int, default=2, help="Maximum rows to evaluate.")
    return parser.parse_args()


def check_python_version() -> None:
    if sys.version_info < (3, 10):
        raise RuntimeError(
            f"Ragas 0.4.3 requires Python 3.10+. Current Python is {sys.version.split()[0]}."
        )


def install_ragas_043_vertexai_compatibility_shims() -> None:
    """Work around obsolete optional VertexAI imports in Ragas 0.4.3.

    The evaluation in this script uses an OpenAI-compatible local endpoint and
    never instantiates these placeholder classes. They only allow Ragas to
    finish importing when a newer langchain-community removed the old modules.
    """

    optional_modules = {
        "langchain_community.chat_models.vertexai": "ChatVertexAI",
        "langchain_community.llms.vertexai": "VertexAI",
    }
    for module_name, class_name in optional_modules.items():
        try:
            importlib.import_module(module_name)
            continue
        except ModuleNotFoundError as exc:
            if exc.name != module_name:
                raise

        module = types.ModuleType(module_name)
        placeholder = type(class_name, (), {})
        setattr(module, class_name, placeholder)
        sys.modules[module_name] = module


def load_config(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(
            f"Config not found: {path}. Copy ragas_config.example.json to ragas_config.json first."
        )
    return json.loads(path.read_text(encoding="utf-8-sig"))


def normalize_base_url(endpoint: str) -> str:
    url = endpoint.rstrip("/")
    suffix = "/chat/completions"
    if url.endswith(suffix):
        url = url[: -len(suffix)]
    return url


def parse_contexts(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]

    text = str(value).strip()
    if not text:
        return []

    for parser in (json.loads, ast.literal_eval):
        try:
            parsed = parser(text)
        except (ValueError, SyntaxError, json.JSONDecodeError):
            continue
        if isinstance(parsed, list):
            return [str(item).strip() for item in parsed if str(item).strip()]
        if isinstance(parsed, str) and parsed.strip():
            return [parsed.strip()]

    return [text]


def parse_reference(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(item).strip() for item in value if str(item).strip())

    text = str(value).strip()
    if not text:
        return ""

    for parser in (json.loads, ast.literal_eval):
        try:
            parsed = parser(text)
        except (ValueError, SyntaxError, json.JSONDecodeError):
            continue
        if isinstance(parsed, list):
            return "\n".join(str(item).strip() for item in parsed if str(item).strip())
        if isinstance(parsed, str):
            return parsed.strip()
    return text


def get_headers(sheet: Any) -> dict[str, int]:
    headers: dict[str, int] = {}
    for column_index, cell in enumerate(sheet[1], start=1):
        if cell.value is not None:
            headers[str(cell.value).strip()] = column_index
    return headers


def ensure_output_columns(sheet: Any, headers: dict[str, int]) -> dict[str, int]:
    next_column = sheet.max_column + 1
    for name in METRIC_COLUMNS:
        if name not in headers:
            sheet.cell(row=1, column=next_column, value=name)
            headers[name] = next_column
            next_column += 1
    return headers


def cell_text(sheet: Any, row: int, column: int) -> str:
    value = sheet.cell(row=row, column=column).value
    return "" if value is None else str(value).strip()


def result_value(result: Any) -> float:
    value = getattr(result, "value", result)
    return float(value)


async def score_row(
    metrics: dict[str, Any],
    question: str,
    answer: str,
    contexts: list[str],
    reference: str,
) -> tuple[dict[str, Optional[float]], list[str]]:
    scores: dict[str, Optional[float]] = {}
    errors: list[str] = []

    metric_inputs = {
        "context_precision": {
            "user_input": question,
            "retrieved_contexts": contexts,
            "reference": reference,
        },
        "context_recall": {
            "user_input": question,
            "retrieved_contexts": contexts,
            "reference": reference,
        },
        "faithfulness": {
            "user_input": question,
            "response": answer,
            "retrieved_contexts": contexts,
        },
        "factual_correctness": {
            "response": answer,
            "reference": reference,
        },
    }

    for name, metric in metrics.items():
        try:
            result = await metric.ascore(**metric_inputs[name])
            scores[name] = result_value(result)
        except Exception as exc:  # noqa: BLE001
            scores[name] = None
            errors.append(f"{name}: {exc}")

    return scores, errors


def write_summary(workbook: Any, sheet: Any, headers: dict[str, int], evaluated_rows: list[int]) -> None:
    summary_name = "ragas_summary"
    if summary_name in workbook.sheetnames:
        del workbook[summary_name]
    summary = workbook.create_sheet(summary_name)
    summary.append(["metric", "mean", "valid_count"])

    for metric_name in METRIC_COLUMNS[:-1]:
        values = []
        column = headers[metric_name]
        for row in evaluated_rows:
            value = sheet.cell(row=row, column=column).value
            if isinstance(value, (int, float)) and not math.isnan(float(value)):
                values.append(float(value))
        mean = sum(values) / len(values) if values else None
        summary.append([metric_name, mean, len(values)])


async def main_async() -> int:
    args = parse_args()
    check_python_version()
    config = load_config(args.config)
    install_ragas_043_vertexai_compatibility_shims()

    try:
        from openai import AsyncOpenAI
        from ragas.llms import llm_factory
        from ragas.metrics.collections import (
            ContextPrecision,
            ContextRecall,
            FactualCorrectness,
            Faithfulness,
        )
    except ImportError as exc:
        raise RuntimeError("Missing dependencies. Run: pip install -r requirements.txt") from exc

    judge = config["judge"]
    client = AsyncOpenAI(
        api_key=judge.get("api_key") or "not-required",
        base_url=normalize_base_url(judge["endpoint"]),
        timeout=float(judge.get("timeout", 180)),
        max_retries=int(judge.get("retries", 2)),
    )
    llm = llm_factory(judge["model"], client=client)
    metrics = {
        "context_precision": ContextPrecision(llm=llm),
        "context_recall": ContextRecall(llm=llm),
        "faithfulness": Faithfulness(llm=llm),
        "factual_correctness": FactualCorrectness(llm=llm),
    }

    workbook = load_workbook(args.input)
    sheet = workbook[args.sheet] if args.sheet else workbook.active
    headers = get_headers(sheet)

    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ValueError(f"Missing Excel columns: {', '.join(missing)}")

    headers = ensure_output_columns(sheet, headers)
    args.output.parent.mkdir(parents=True, exist_ok=True)

    evaluated_rows: list[int] = []
    processed = 0
    for row in range(2, sheet.max_row + 1):
        if processed >= args.limit:
            break

        question = cell_text(sheet, row, headers["question"])
        answer = cell_text(sheet, row, headers["answer"])
        contexts = parse_contexts(sheet.cell(row=row, column=headers["contexts"]).value)
        reference = parse_reference(sheet.cell(row=row, column=headers["ground_truths"]).value)

        if not any((question, answer, contexts, reference)):
            continue

        validation_errors = []
        if not question:
            validation_errors.append("question is empty")
        if not answer:
            validation_errors.append("answer is empty")
        if not contexts:
            validation_errors.append("contexts is empty")
        if not reference:
            validation_errors.append("ground_truths is empty")

        print(f"[{processed + 1}/{args.limit}] Evaluating Excel row {row}", flush=True)
        if validation_errors:
            sheet.cell(row=row, column=headers["ragas_error"], value="; ".join(validation_errors))
        else:
            scores, errors = await score_row(metrics, question, answer, contexts, reference)
            for metric_name, value in scores.items():
                sheet.cell(row=row, column=headers[metric_name], value=value)
            sheet.cell(row=row, column=headers["ragas_error"], value="\n".join(errors))

        evaluated_rows.append(row)
        processed += 1
        write_summary(workbook, sheet, headers, evaluated_rows)
        workbook.save(args.output)

    print(f"Evaluated {processed} row(s). Result: {args.output.resolve()}", flush=True)
    return 0


def main() -> int:
    try:
        return asyncio.run(main_async())
    except KeyboardInterrupt:
        print("Interrupted.", file=sys.stderr)
        return 130
    except Exception as exc:  # noqa: BLE001
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
