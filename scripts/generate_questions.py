from __future__ import annotations

import argparse
import json
import math
import re
import sys
import time
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_CONFIG: dict[str, Any] = {
    "vlm": {
        "endpoint": "http://localhost:8000/v1/chat/completions",
        "api_key": "your_api_key",
        "model": "Deepseek-V4-flash",
        "max_tokens": 8192,
        "timeout": 120,
        "retries": 2,
    },
    "generation": {
        "generation_mode": "section_detailed",
        "target_count": 30,
        "candidate_multiplier": 1.2,
        "cross_doc_count": 2,
        "language": "简体中文",
        "max_chunk_chars": 80000,
        "chunk_overlap_chars": 1000,
        "max_chunks_per_doc": 5,
        "max_total_chunks": 30,
        "min_chunk_chars": 800,
        "max_llm_input_chars": 180000,
        "knowledge_points_per_chunk": 6,
        "questions_per_doc_min": 3,
        "questions_per_doc_max": 35,
        "questions_per_chunk": 2,
        "answer_min_paragraphs": 2,
        "answer_max_paragraphs": 4,
        "answer_min_chars": 350,
        "answer_max_chars": 1200,
        "max_candidates_total": 45,
        "continue_on_error": True,
        "max_consecutive_llm_errors": 3,
        "checkpoint_enabled": True,
        "checkpoint_interval_chunks": 1,
        "checkpoint_json": True,
        "json_repair_retries": 1,
        "temperature": 0.2,
        "section_include_keywords": [
            "功能",
            "设计",
            "架构",
            "流程",
            "接口",
            "模块",
            "软件层",
            "CPD",
            "处理",
            "异常",
            "状态",
            "数据",
            "时序",
            "配置",
            "交互",
            "约束"
        ],
        "section_exclude_keywords": [
            "目录",
            "修订记录",
            "变更记录",
            "术语",
            "缩略语",
            "参考文献",
            "参考资料",
            "引言"
        ],
    },
}


QUESTION_TYPES = [
    "fact",
    "process",
    "constraint",
    "interface",
    "exception",
    "dependency",
    "cross_doc",
    "section_explanation",
    "end_to_end_process",
    "module_behavior",
    "architecture",
]


@dataclass
class Document:
    doc_id: str
    doc_name: str
    path: str
    content: str


@dataclass
class Chunk:
    chunk_id: str
    doc_id: str
    doc_name: str
    source_path: str
    section_title: str
    heading_path: str
    start_char: int
    end_char: int
    content: str


class LLMClient:
    def __init__(self, config: dict[str, Any]) -> None:
        self.endpoint = config["endpoint"]
        self.api_key = config.get("api_key", "")
        self.model = config["model"]
        self.max_tokens = config.get("max_tokens")
        self.timeout = int(config.get("timeout", 120))
        self.retries = int(config.get("retries", 2))

    def chat(self, messages: list[dict[str, str]], temperature: float = 0.2) -> str:
        headers = {"Content-Type": "application/json"}
        if self.api_key:
            headers["Authorization"] = f"Bearer {self.api_key}"

        payload = {
            "model": self.model,
            "messages": messages,
            "temperature": temperature,
        }
        if self.max_tokens is not None:
            payload["max_tokens"] = int(self.max_tokens)

        last_error: Exception | None = None
        for attempt in range(self.retries + 1):
            try:
                response = requests.post(
                    self.endpoint,
                    headers=headers,
                    json=payload,
                    timeout=self.timeout,
                )
                response.raise_for_status()
                data = response.json()
                return data["choices"][0]["message"]["content"]
            except Exception as exc:  # noqa: BLE001
                last_error = exc
                if attempt < self.retries:
                    time.sleep(1.5 * (attempt + 1))
        raise RuntimeError(f"LLM request failed after retries: {last_error}") from last_error


def deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    result = dict(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(result.get(key), dict):
            result[key] = deep_merge(result[key], value)
        else:
            result[key] = value
    return result


def load_config(path: Path | None) -> dict[str, Any]:
    if path is None:
        return DEFAULT_CONFIG
    with path.open("r", encoding="utf-8") as f:
        user_config = json.load(f)
    return deep_merge(DEFAULT_CONFIG, user_config)


def read_markdown_files(input_dir: Path | None, input_files: list[Path] | None) -> list[Document]:
    paths: list[Path] = []
    if input_dir:
        paths.extend(sorted(input_dir.rglob("*.md")))
        paths.extend(sorted(input_dir.rglob("*.markdown")))
    if input_files:
        paths.extend(input_files)

    unique_paths = []
    seen = set()
    for path in paths:
        resolved = path.resolve()
        if resolved not in seen:
            seen.add(resolved)
            unique_paths.append(resolved)

    documents: list[Document] = []
    for idx, path in enumerate(unique_paths, start=1):
        if not path.exists():
            raise FileNotFoundError(f"Markdown file not found: {path}")
        content = path.read_text(encoding="utf-8-sig", errors="replace")
        if content.strip():
            documents.append(
                Document(
                    doc_id=f"DOC_{idx:03d}",
                    doc_name=path.stem,
                    path=str(path),
                    content=normalize_markdown(content),
                )
            )
    if not documents:
        raise ValueError("No non-empty markdown files found.")
    return documents


def normalize_markdown(content: str) -> str:
    content = content.replace("\r\n", "\n").replace("\r", "\n")
    content = re.sub(r"\n{4,}", "\n\n\n", content)
    return content.strip()


def split_document(document: Document, max_chars: int, overlap_chars: int) -> list[Chunk]:
    sections = split_by_headings(document.content)
    chunks: list[Chunk] = []

    for section_title, heading_path, start, text in sections:
        if len(text) <= max_chars:
            chunks.append(
                build_chunk(document, len(chunks) + 1, section_title, heading_path, start, start + len(text), text)
            )
            continue

        local_start = 0
        while local_start < len(text):
            local_end = min(local_start + max_chars, len(text))
            if local_end < len(text):
                paragraph_break = text.rfind("\n\n", local_start, local_end)
                if paragraph_break > local_start + int(max_chars * 0.55):
                    local_end = paragraph_break
            piece = text[local_start:local_end].strip()
            if piece:
                chunks.append(
                    build_chunk(
                        document,
                        len(chunks) + 1,
                        section_title,
                        heading_path,
                        start + local_start,
                        start + local_end,
                        piece,
                    )
                )
            if local_end >= len(text):
                break
            local_start = max(local_end - overlap_chars, local_start + 1)

    return chunks


def limit_chunks_for_fast_run(chunks: list[Chunk], generation_config: dict[str, Any]) -> list[Chunk]:
    max_per_doc = int(generation_config.get("max_chunks_per_doc", 0) or 0)
    max_total = int(generation_config.get("max_total_chunks", 0) or 0)
    selection_mode = str(generation_config.get("chunk_selection", "scored")).lower()

    selected: list[Chunk] = []
    for _, doc_chunks in group_by(chunks, "doc_id").items():
        if selection_mode == "even":
            selected.extend(select_evenly(doc_chunks, max_per_doc))
        else:
            selected.extend(select_high_value_chunks(doc_chunks, max_per_doc, generation_config))

    if max_total > 0:
        if selection_mode == "even":
            selected = select_evenly(selected, max_total)
        else:
            selected = select_high_value_chunks(selected, max_total, generation_config)
    return selected


def select_high_value_chunks(chunks: list[Chunk], limit: int, generation_config: dict[str, Any]) -> list[Chunk]:
    if limit <= 0 or len(chunks) <= limit:
        return chunks

    scored = [(score_chunk(chunk, generation_config), index, chunk) for index, chunk in enumerate(chunks)]
    scored.sort(key=lambda item: (-item[0], item[1]))
    selected = [chunk for score, _, chunk in scored if score > 0][:limit]

    if len(selected) < limit:
        selected_ids = {chunk.chunk_id for chunk in selected}
        for _, _, chunk in scored:
            if chunk.chunk_id not in selected_ids:
                selected.append(chunk)
                selected_ids.add(chunk.chunk_id)
            if len(selected) >= limit:
                break

    return sorted(selected, key=lambda chunk: (chunk.doc_id, chunk.start_char))


def score_chunk(chunk: Chunk, generation_config: dict[str, Any]) -> int:
    title_text = f"{chunk.section_title} {chunk.heading_path}"
    content_text = chunk.content
    combined = f"{title_text}\n{content_text[:5000]}"
    min_chars = int(generation_config.get("min_chunk_chars", 0) or 0)
    include_keywords = ensure_list(generation_config.get("section_include_keywords"))
    exclude_keywords = ensure_list(generation_config.get("section_exclude_keywords"))

    if len(content_text.strip()) < min_chars:
        return 0

    if any(keyword and keyword.lower() in title_text.lower() for keyword in exclude_keywords):
        return 0

    score = 1
    if len(content_text) >= 3000:
        score += 2
    if len(content_text) >= 10000:
        score += 2
    if len(content_text) >= 30000:
        score += 1

    for keyword in include_keywords:
        if not keyword:
            continue
        keyword_lower = keyword.lower()
        if keyword_lower in title_text.lower():
            score += 8
        if keyword_lower in combined.lower():
            score += 2

    structural_patterns = [
        r"\|.+\|.+\|",
        r"接口",
        r"输入",
        r"输出",
        r"异常",
        r"状态",
        r"步骤",
        r"流程",
        r"调用",
    ]
    for pattern in structural_patterns:
        if re.search(pattern, combined, re.IGNORECASE):
            score += 2

    return score


def select_evenly(items: list[Any], limit: int) -> list[Any]:
    if limit <= 0 or len(items) <= limit:
        return items
    if limit == 1:
        return [items[0]]

    step = (len(items) - 1) / (limit - 1)
    indexes = sorted({round(i * step) for i in range(limit)})
    selected = [items[index] for index in indexes]

    cursor = 0
    while len(selected) < limit and cursor < len(items):
        if items[cursor] not in selected:
            selected.append(items[cursor])
        cursor += 1
    return selected[:limit]


def split_by_headings(content: str) -> list[tuple[str, str, int, str]]:
    heading_pattern = re.compile(r"^(#{1,6})\s+(.+?)\s*$", re.MULTILINE)
    matches = list(heading_pattern.finditer(content))
    if not matches:
        return [("全文", "全文", 0, content)]

    sections: list[tuple[str, str, int, str]] = []
    heading_stack: list[tuple[int, str]] = []

    if matches[0].start() > 0:
        preface = content[: matches[0].start()].strip()
        if preface:
            sections.append(("文档前言", "文档前言", 0, preface))

    for idx, match in enumerate(matches):
        level = len(match.group(1))
        title = match.group(2).strip()
        while heading_stack and heading_stack[-1][0] >= level:
            heading_stack.pop()
        heading_stack.append((level, title))

        section_start = match.start()
        section_end = matches[idx + 1].start() if idx + 1 < len(matches) else len(content)
        section_text = content[section_start:section_end].strip()
        heading_path = " > ".join(item[1] for item in heading_stack)
        sections.append((title, heading_path, section_start, section_text))

    return sections


def build_chunk(
    document: Document,
    chunk_index: int,
    section_title: str,
    heading_path: str,
    start_char: int,
    end_char: int,
    content: str,
) -> Chunk:
    return Chunk(
        chunk_id=f"{document.doc_id}_C{chunk_index:03d}",
        doc_id=document.doc_id,
        doc_name=document.doc_name,
        source_path=document.path,
        section_title=section_title,
        heading_path=heading_path,
        start_char=start_char,
        end_char=end_char,
        content=content,
    )


def truncate_text(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[内容因长度限制被截断]"


def safe_print(message: str) -> None:
    try:
        print(message, flush=True)
    except UnicodeEncodeError:
        encoding = getattr(sys.stdout, "encoding", None) or "utf-8"
        safe_message = message.encode(encoding, errors="replace").decode(encoding, errors="replace")
        print(safe_message, flush=True)


def extract_json_array(text: str) -> list[dict[str, Any]]:
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned)

    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\[[\s\S]*\]", cleaned)
        if not match:
            raise
        data = json.loads(match.group(0))

    if not isinstance(data, list):
        raise ValueError("Expected a JSON array from LLM output.")

    normalized = []
    for item in data:
        if isinstance(item, dict):
            normalized.append(item)
    return normalized


def chat_json_array(
    client: LLMClient,
    messages: list[dict[str, str]],
    generation_config: dict[str, Any],
) -> list[dict[str, Any]]:
    repair_retries = int(generation_config.get("json_repair_retries", 1))
    last_output = ""
    last_error: Exception | None = None

    for attempt in range(repair_retries + 1):
        if attempt == 0:
            active_messages = messages
        else:
            active_messages = [
                *messages,
                {
                    "role": "assistant",
                    "content": last_output,
                },
                {
                    "role": "user",
                    "content": (
                        "上一次输出无法被解析为严格 JSON 数组。请只输出修正后的 JSON 数组，"
                        "不要输出 Markdown 代码块、解释文字或额外前后缀。"
                    ),
                },
            ]

        last_output = client.chat(active_messages, temperature=generation_config["temperature"])
        try:
            return extract_json_array(last_output)
        except Exception as exc:  # noqa: BLE001
            last_error = exc

    raise RuntimeError(f"Failed to parse JSON array from LLM output: {last_error}") from last_error


def extract_knowledge_points(
    client: LLMClient,
    chunk: Chunk,
    generation_config: dict[str, Any],
) -> list[dict[str, Any]]:
    system_prompt = (
        "你是一位资深软件研发知识库评测专家，擅长从软件设计说明书中抽取可用于 RAG "
        "评测的事实、流程、约束、接口、异常和依赖关系。所有输出必须严格基于输入文档，"
        "不允许引入外部知识、常识推断或自行补充。只输出 JSON 数组，不要输出解释。"
    )
    user_prompt = f"""
请从以下软件设计说明书片段中抽取可用于出题的知识点。

要求：
1. 每个知识点必须能被原文明确支持。
2. 优先抽取模块职责、业务流程、状态流转、接口输入输出、字段含义、异常处理、约束条件、权限规则、数据依赖、跨模块调用关系。
3. 不要抽取过于宽泛的内容，例如“系统用于提升效率”。
4. 每个知识点需要包含证据摘要和来源位置。
5. 如果内容不足以形成问题，请在 is_questionable 中标记 false。
6. 最多输出 {generation_config["knowledge_points_per_chunk"]} 个知识点。
7. 使用{generation_config["language"]}。

文档信息：
doc_id: {chunk.doc_id}
doc_name: {chunk.doc_name}
chunk_id: {chunk.chunk_id}
section_title: {chunk.section_title}
heading_path: {chunk.heading_path}

输出 JSON 数组，每个对象字段为：
knowledge_id, doc_id, doc_name, chunk_id, topic, knowledge_type, fact, evidence_summary,
source_section, source_chunk_id, is_questionable

文档内容：
{truncate_text(chunk.content, generation_config["max_llm_input_chars"])}
""".strip()

    items = chat_json_array(
        client,
        [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
        generation_config,
    )
    result = []
    max_points = int(generation_config["knowledge_points_per_chunk"])
    for idx, item in enumerate(items[:max_points], start=1):
        item.setdefault("knowledge_id", f"{chunk.chunk_id}_K{idx:03d}")
        item["doc_id"] = chunk.doc_id
        item["doc_name"] = chunk.doc_name
        item["chunk_id"] = chunk.chunk_id
        item.setdefault("source_chunk_id", chunk.chunk_id)
        item.setdefault("source_section", chunk.heading_path)
        result.append(item)
    return result


def generate_doc_candidates(
    client: LLMClient,
    document: Document,
    doc_chunks: list[Chunk],
    knowledge_points: list[dict[str, Any]],
    generation_config: dict[str, Any],
    target_questions: int,
) -> list[dict[str, Any]]:
    if not knowledge_points:
        return []

    q_min = int(generation_config["questions_per_doc_min"])
    q_max = int(generation_config["questions_per_doc_max"])
    target = min(q_max, max(q_min, target_questions))
    source_excerpt = build_source_excerpt(doc_chunks, generation_config["max_llm_input_chars"] // 3)

    system_prompt = (
        "你是一位软件研发知识库 TestBench 出题专家。你需要根据软件设计说明书中的知识点，"
        "生成适合评估 RAG 知识库效果的问题和标准答案。Ground Truth 必须完全由输入知识点"
        "和原文依据支持，不允许补充文档外信息。只输出 JSON 数组，不要输出解释。"
    )
    user_prompt = f"""
请基于以下文档知识点生成候选问答。

出题要求：
1. 必须只生成 {target} 道候选题，不要超过 {target} 道。
2. 问题必须具体，避免“介绍一下某模块”这类宽泛问题。
3. 每道题都必须可以从输入文档中直接或间接回答。
4. Ground Truth 应完整回答问题，语言清晰、无歧义。
5. Ground Truth 不要大段复制原文，但必须忠实于原文。
6. 题型覆盖 fact、process、constraint、interface、exception、dependency。
7. 每道题必须给出 required_evidence，用于后续判断检索 chunk 是否召回关键信息。
8. 每道题必须给出 source_chunk_ids 和 related_knowledge_ids。
9. 使用{generation_config["language"]}。

文档信息：
doc_id: {document.doc_id}
doc_name: {document.doc_name}

知识点：
{json.dumps(knowledge_points, ensure_ascii=False)}

原文摘录：
{source_excerpt}

输出 JSON 数组，每个对象字段为：
candidate_id, question, ground_truth, question_type, difficulty, answer_scope,
source_doc, source_sections, source_chunk_ids, related_knowledge_ids, required_evidence,
evidence_summary, quality_reason
""".strip()

    items = chat_json_array(
        client,
        [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
        generation_config,
    )
    result = []
    for idx, item in enumerate(items[:target], start=1):
        item.setdefault("candidate_id", f"{document.doc_id}_Q{idx:03d}")
        item.setdefault("source_doc", document.doc_name)
        item.setdefault("answer_scope", "single_doc")
        result.append(normalize_question_item(item))
    return result


def generate_cross_doc_candidates(
    client: LLMClient,
    knowledge_points: list[dict[str, Any]],
    generation_config: dict[str, Any],
) -> list[dict[str, Any]]:
    count = int(generation_config.get("cross_doc_count", 0))
    if count <= 0 or len({kp.get("doc_id") for kp in knowledge_points}) < 2:
        return []

    compact_points = compact_knowledge_points(knowledge_points, generation_config["max_llm_input_chars"])
    system_prompt = (
        "你是一位严格的软件研发知识库综合题出题专家。你需要生成少量跨文档或跨模块问题，"
        "每道题至少关联两个知识点。问题和答案必须完全由输入知识点支持。只输出 JSON 数组。"
    )
    user_prompt = f"""
请基于以下多个文档的知识点，生成 {count} 道跨文档/跨模块候选问答。

要求：
1. 每道题至少关联 2 个 related_knowledge_ids。
2. 问题必须能被给定知识点完全支持。
3. 不要生成需要业务猜测或外部常识的问题。
4. Ground Truth 要说明不同模块、流程或数据之间的关系。
5. answer_scope 标记为 cross_doc 或 cross_section。
6. 使用{generation_config["language"]}。

知识点集合：
{compact_points}

输出 JSON 数组，每个对象字段为：
candidate_id, question, ground_truth, question_type, difficulty, answer_scope,
source_doc, source_sections, source_chunk_ids, related_knowledge_ids, required_evidence,
evidence_summary, quality_reason
""".strip()

    items = chat_json_array(
        client,
        [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
        generation_config,
    )
    result = []
    for idx, item in enumerate(items[:count], start=1):
        item.setdefault("candidate_id", f"CROSS_Q{idx:03d}")
        item.setdefault("answer_scope", "cross_doc")
        item.setdefault("question_type", "cross_doc")
        result.append(normalize_question_item(item))
    return result


def generate_section_detailed_candidates(
    client: LLMClient,
    chunk: Chunk,
    generation_config: dict[str, Any],
    target_questions: int,
) -> list[dict[str, Any]]:
    if target_questions <= 0:
        return []

    answer_min_paragraphs = int(generation_config.get("answer_min_paragraphs", 2))
    answer_max_paragraphs = int(generation_config.get("answer_max_paragraphs", 4))
    answer_min_chars = int(generation_config.get("answer_min_chars", 350))
    answer_max_chars = int(generation_config.get("answer_max_chars", 1200))

    system_prompt = (
        "你是一位软件研发知识库 TestBench 出题专家。你需要基于软件设计文档的一个完整章节，"
        "生成适合评估 RAG 知识库效果的综合型问题和标准答案。标准答案必须完全基于输入章节，"
        "不允许补充文档外信息。只输出 JSON 数组，不要输出解释。"
    )
    user_prompt = f"""
请基于以下软件设计文档章节生成 {target_questions} 道综合型问答。

问题要求：
1. 必须只生成 {target_questions} 道题，不要超过 {target_questions} 道。
2. 问题要围绕完整功能、CPD 项、软件层架构、模块职责、端到端流程、接口协作、异常处理或关键约束展开。
3. 问题应适合研发、测试、维护人员查阅设计文档时提出。
4. 避免只能一句话回答的问题，例如“某字段是什么”“某接口叫什么”“是否支持某功能”。
5. 如果当前章节不足以支撑多段答案，可以少生成或返回空数组。

Ground Truth 要求：
1. 不要只写一句话；每个 ground_truth 写成 {answer_min_paragraphs}~{answer_max_paragraphs} 个自然段。
2. 总长度建议控制在 {answer_min_chars}~{answer_max_chars} 个中文字符。
3. 回答应先概括该章节描述的功能/架构/流程目标，再展开说明关键设计内容。
4. 如果章节中包含流程、接口、字段、状态、约束、异常、上下游依赖或软件层交互，应在答案中一并说明。
5. 必须忠实于输入章节；文档没有明确说明的内容不要推断。
6. 使用{generation_config["language"]}。

文档信息：
doc_id: {chunk.doc_id}
doc_name: {chunk.doc_name}
chunk_id: {chunk.chunk_id}
section_title: {chunk.section_title}
heading_path: {chunk.heading_path}

章节内容：
{truncate_text(chunk.content, generation_config["max_llm_input_chars"])}

输出 JSON 数组，每个对象字段为：
candidate_id, question, ground_truth, question_type, difficulty, answer_scope,
source_doc, source_sections, source_chunk_ids, related_knowledge_ids, required_evidence,
evidence_summary, quality_reason
""".strip()

    items = chat_json_array(
        client,
        [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
        generation_config,
    )
    result = []
    for idx, item in enumerate(items[:target_questions], start=1):
        item.setdefault("candidate_id", f"{chunk.chunk_id}_Q{idx:03d}")
        item.setdefault("source_doc", chunk.doc_name)
        item.setdefault("source_sections", [chunk.heading_path])
        item.setdefault("source_chunk_ids", [chunk.chunk_id])
        item.setdefault("answer_scope", "single_section")
        item.setdefault("question_type", "section_explanation")
        result.append(normalize_question_item(item))
    return result


def review_and_select_questions(
    client: LLMClient,
    candidates: list[dict[str, Any]],
    knowledge_points: list[dict[str, Any]],
    generation_config: dict[str, Any],
) -> list[dict[str, Any]]:
    target_count = int(generation_config["target_count"])
    compact_candidates = truncate_json(candidates, generation_config["max_llm_input_chars"] // 2)
    compact_points = compact_knowledge_points(knowledge_points, generation_config["max_llm_input_chars"] // 2)

    system_prompt = (
        "你是一位严格的软件知识库评测集审核专家。你的任务是从候选问题中筛选最终 TestBench "
        "题集，并修正少量表述问题。你需要优先保证可回答性、证据充分性和题型覆盖。"
        "只输出 JSON 数组，不要输出解释。"
    )
    user_prompt = f"""
请从以下候选问答中筛选最终 TestBench 题集。

筛选目标：
1. 总数控制在 {target_count} 道，如果候选不足可少于该数量。
2. 尽量覆盖所有文档。
3. 题型尽量均衡：fact、process、constraint、interface、exception、dependency、cross_doc。
4. 保留少量跨文档/跨模块问题，但不要超过 20%。
5. 优先选择证据明确、答案边界清晰、人工复核价值高的问题。
6. 在 section_detailed 模式下，优先保留能够支撑多段解释的问题；不要因为 Ground Truth 较长就判定为冗余。
7. 删除重复题、过宽题、答案过短题、不可回答题。
8. Ground Truth 必须完全由证据支持，必要时可轻微修订；如果答案只有一句话，应优先淘汰。
9. 使用{generation_config["language"]}。

候选题：
{compact_candidates}

相关知识点：
{compact_points}

输出 JSON 数组，每个对象字段为：
question_id, question, ground_truth, question_type, difficulty, answer_scope,
source_doc, source_sections, source_chunk_ids, related_knowledge_ids, required_evidence,
evidence_summary, quality_score, review_reason
""".strip()

    selected = chat_json_array(
        client,
        [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
        generation_config,
    )
    final_questions = []
    for idx, item in enumerate(selected[:target_count], start=1):
        item = normalize_question_item(item)
        item["question_id"] = item.get("question_id") or f"Q{idx:03d}"
        final_questions.append(item)
    return final_questions


def build_source_excerpt(chunks: list[Chunk], max_chars: int) -> str:
    pieces = []
    used = 0
    for chunk in chunks:
        header = f"\n\n[{chunk.chunk_id} | {chunk.heading_path}]\n"
        available = max_chars - used - len(header)
        if available <= 0:
            break
        text = chunk.content[:available]
        pieces.append(header + text)
        used += len(header) + len(text)
    return "".join(pieces).strip()


def compact_knowledge_points(knowledge_points: list[dict[str, Any]], max_chars: int) -> str:
    compact = []
    for kp in knowledge_points:
        compact.append(
            {
                "knowledge_id": kp.get("knowledge_id"),
                "doc_id": kp.get("doc_id"),
                "doc_name": kp.get("doc_name"),
                "chunk_id": kp.get("chunk_id") or kp.get("source_chunk_id"),
                "topic": kp.get("topic"),
                "knowledge_type": kp.get("knowledge_type"),
                "fact": kp.get("fact"),
                "evidence_summary": kp.get("evidence_summary"),
                "source_section": kp.get("source_section"),
            }
        )
    return truncate_json(compact, max_chars)


def truncate_json(data: Any, max_chars: int) -> str:
    text = json.dumps(data, ensure_ascii=False, indent=2)
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + '\n"...内容因长度限制被截断"'


def normalize_question_item(item: dict[str, Any]) -> dict[str, Any]:
    normalized = dict(item)
    normalized["question_type"] = normalize_enum(normalized.get("question_type"), QUESTION_TYPES, "fact")
    normalized["difficulty"] = normalize_enum(normalized.get("difficulty"), ["easy", "medium", "hard"], "medium")
    normalized["answer_scope"] = normalize_enum(
        normalized.get("answer_scope"),
        ["single_section", "single_doc", "cross_section", "cross_doc"],
        "single_doc",
    )
    for key in ["source_sections", "source_chunk_ids", "related_knowledge_ids", "required_evidence"]:
        normalized[key] = ensure_list(normalized.get(key))
    return normalized


def normalize_enum(value: Any, allowed: list[str], default: str) -> str:
    if isinstance(value, str):
        value = value.strip()
        if value in allowed:
            return value
    return default


def ensure_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item) for item in value if item is not None]
    if isinstance(value, str):
        if not value.strip():
            return []
        return [value.strip()]
    return [str(value)]


def write_excel(
    output_path: Path,
    final_questions: list[dict[str, Any]],
    candidates: list[dict[str, Any]],
    knowledge_points: list[dict[str, Any]],
    chunks: list[Chunk],
    config: dict[str, Any],
    errors: list[dict[str, Any]] | None = None,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    add_sheet(workbook, "final_questions", final_questions)
    add_sheet(workbook, "candidates", candidates)
    add_sheet(workbook, "knowledge_points", knowledge_points)
    add_sheet(workbook, "chunks", [chunk_to_row(chunk) for chunk in chunks])
    add_sheet(workbook, "errors", errors or [])
    add_sheet(
        workbook,
        "run_config",
        [
            {"key": key, "value": json.dumps(value, ensure_ascii=False)}
            for key, value in flatten_top_level_config(config).items()
        ],
    )

    workbook.save(output_path)


def save_progress(
    output_path: Path,
    stage: str,
    final_questions: list[dict[str, Any]],
    candidates: list[dict[str, Any]],
    knowledge_points: list[dict[str, Any]],
    chunks: list[Chunk],
    config: dict[str, Any],
    errors: list[dict[str, Any]],
) -> None:
    if not bool(config["generation"].get("checkpoint_enabled", True)):
        return

    write_excel(output_path, final_questions, candidates, knowledge_points, chunks, config, errors)

    if bool(config["generation"].get("checkpoint_json", True)):
        checkpoint_path = output_path.with_suffix(".checkpoint.json")
        checkpoint = {
            "stage": stage,
            "saved_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "final_questions": final_questions,
            "candidates": candidates,
            "knowledge_points": knowledge_points,
            "chunks": [chunk_to_row(chunk) for chunk in chunks],
            "errors": errors,
            "config": config,
        }
        checkpoint_path.parent.mkdir(parents=True, exist_ok=True)
        checkpoint_path.write_text(json.dumps(checkpoint, ensure_ascii=False, indent=2), encoding="utf-8")


def add_error(errors: list[dict[str, Any]], stage: str, item_id: str, message: str) -> None:
    errors.append(
        {
            "stage": stage,
            "item_id": item_id,
            "message": message,
            "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        }
    )


def fallback_select_questions(candidates: list[dict[str, Any]], target_count: int) -> list[dict[str, Any]]:
    final_questions = []
    seen_questions = set()
    for item in candidates:
        question = str(item.get("question", "")).strip()
        ground_truth = str(item.get("ground_truth", "")).strip()
        if not question or not ground_truth or question in seen_questions:
            continue
        seen_questions.add(question)
        final_item = normalize_question_item(dict(item))
        final_item["question_id"] = f"Q{len(final_questions) + 1:03d}"
        final_item.setdefault("quality_score", "")
        final_item.setdefault("review_reason", "LLM 终筛失败或跳过，使用候选题兜底入选。")
        final_questions.append(final_item)
        if len(final_questions) >= target_count:
            break
    return final_questions


def chunk_to_row(chunk: Chunk) -> dict[str, Any]:
    row = asdict(chunk)
    row["content_preview"] = chunk.content[:1000]
    row.pop("content", None)
    return row


def flatten_top_level_config(config: dict[str, Any]) -> dict[str, Any]:
    flat = {}
    for section, values in config.items():
        if isinstance(values, dict):
            for key, value in values.items():
                flat[f"{section}.{key}"] = value
        else:
            flat[section] = values
    return flat


def add_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title=title)
    if not rows:
        sheet.append(["empty"])
        return

    headers = collect_headers(rows)
    sheet.append(headers)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        sheet.append([format_cell_value(row.get(header)) for header in headers])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for col_idx, header in enumerate(headers, start=1):
        values = [str(header)] + [str(format_cell_value(row.get(header))) for row in rows[:100]]
        max_len = min(max(len(value) for value in values), 80)
        sheet.column_dimensions[get_column_letter(col_idx)].width = max(12, max_len + 2)

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def collect_headers(rows: list[dict[str, Any]]) -> list[str]:
    preferred_order = [
        "question_id",
        "candidate_id",
        "question",
        "ground_truth",
        "question_type",
        "difficulty",
        "answer_scope",
        "source_doc",
        "source_sections",
        "source_chunk_ids",
        "related_knowledge_ids",
        "required_evidence",
        "evidence_summary",
        "quality_score",
        "review_reason",
        "quality_reason",
        "knowledge_id",
        "doc_id",
        "doc_name",
        "chunk_id",
        "topic",
        "knowledge_type",
        "fact",
        "source_section",
        "is_questionable",
        "source_path",
        "heading_path",
        "section_title",
        "start_char",
        "end_char",
        "content_preview",
    ]
    seen = set()
    headers = []
    for header in preferred_order:
        if any(header in row for row in rows):
            headers.append(header)
            seen.add(header)
    for row in rows:
        for key in row:
            if key not in seen:
                headers.append(key)
                seen.add(key)
    return headers


def format_cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, indent=2)
    return str(value)


def parse_args(argv: Iterable[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate RAG TestBench questions from markdown design docs.")
    input_group = parser.add_mutually_exclusive_group(required=True)
    input_group.add_argument("--input-dir", type=Path, help="Directory containing markdown files.")
    input_group.add_argument("--input-files", nargs="+", type=Path, help="Explicit markdown files.")
    parser.add_argument("--config", type=Path, default=None, help="JSON config path.")
    parser.add_argument("--output", type=Path, default=Path("output/questions_ground_truth.xlsx"), help="Output xlsx path.")
    return parser.parse_args(list(argv))


def main(argv: Iterable[str]) -> int:
    args = parse_args(argv)
    config = load_config(args.config)
    gen_config = config["generation"]
    continue_on_error = bool(gen_config.get("continue_on_error", True))
    target_count = int(gen_config["target_count"])
    candidate_budget = min(
        int(gen_config.get("max_candidates_total", 45)),
        max(target_count, math.ceil(target_count * float(gen_config.get("candidate_multiplier", 1.2)))),
    )

    documents = read_markdown_files(args.input_dir, args.input_files)
    raw_chunks: list[Chunk] = []
    for document in documents:
        raw_chunks.extend(
            split_document(
                document,
                max_chars=int(gen_config["max_chunk_chars"]),
                overlap_chars=int(gen_config["chunk_overlap_chars"]),
            )
        )
    all_chunks = limit_chunks_for_fast_run(raw_chunks, gen_config)

    safe_print(
        f"Loaded {len(documents)} documents, split into {len(raw_chunks)} chunks, "
        f"selected {len(all_chunks)} chunks for this run."
    )

    client = LLMClient(config["vlm"])

    errors: list[dict[str, Any]] = []
    candidates: list[dict[str, Any]] = []
    final_questions: list[dict[str, Any]] = []
    all_knowledge_points: list[dict[str, Any]] = []
    save_progress(args.output, "chunks_selected", final_questions, candidates, all_knowledge_points, all_chunks, config, errors)

    generation_mode = str(gen_config.get("generation_mode", "section_detailed")).lower()
    if generation_mode == "section_detailed":
        questions_per_chunk = max(1, int(gen_config.get("questions_per_chunk", 2)))
        max_consecutive_errors = int(gen_config.get("max_consecutive_llm_errors", 3))
        consecutive_errors = 0

        for idx, chunk in enumerate(all_chunks, start=1):
            if len(candidates) >= candidate_budget:
                break
            remaining_budget = max(0, candidate_budget - len(candidates))
            target_for_chunk = min(questions_per_chunk, remaining_budget)
            if target_for_chunk <= 0:
                break

            safe_print(
                f"[{idx}/{len(all_chunks)}] Generating detailed questions: "
                f"{chunk.chunk_id} {chunk.heading_path}"
            )
            try:
                chunk_candidates = generate_section_detailed_candidates(
                    client,
                    chunk,
                    gen_config,
                    target_for_chunk,
                )
                candidates.extend(chunk_candidates[:remaining_budget])
                consecutive_errors = 0
            except Exception as exc:  # noqa: BLE001
                consecutive_errors += 1
                add_error(errors, "generate_section_detailed_candidates", chunk.chunk_id, str(exc))
                safe_print(f"  Failed, skipped: {exc}")
                if not continue_on_error:
                    raise
                if max_consecutive_errors > 0 and consecutive_errors >= max_consecutive_errors:
                    add_error(
                        errors,
                        "generate_section_detailed_candidates",
                        "abort_remaining_chunks",
                        f"Stopped after {consecutive_errors} consecutive LLM errors.",
                    )
                    safe_print(f"  Stopped remaining chunks after {consecutive_errors} consecutive LLM errors.")
                    break

            save_progress(
                args.output,
                f"detailed_candidates_chunk_{idx}",
                final_questions,
                candidates,
                all_knowledge_points,
                all_chunks,
                config,
                errors,
            )

        safe_print(f"Generated {len(candidates)} candidates.")
        safe_print("Reviewing and selecting final questions")
        if candidates:
            try:
                final_questions = review_and_select_questions(client, candidates, all_knowledge_points, gen_config)
            except Exception as exc:  # noqa: BLE001
                add_error(errors, "review_and_select_questions", "final_review", str(exc))
                safe_print(f"  Final review failed, using fallback selection: {exc}")
                if not continue_on_error:
                    raise
                final_questions = fallback_select_questions(candidates, target_count)
        else:
            add_error(errors, "review_and_select_questions", "final_review", "No candidates generated.")

        if not final_questions and candidates:
            final_questions = fallback_select_questions(candidates, target_count)

        safe_print(f"Selected {len(final_questions)} final questions.")
        write_excel(args.output, final_questions, candidates, all_knowledge_points, all_chunks, config, errors)
        save_progress(args.output, "completed", final_questions, candidates, all_knowledge_points, all_chunks, config, errors)
        safe_print(f"Wrote Excel: {args.output.resolve()}")
        return 0

    checkpoint_interval = max(1, int(gen_config.get("checkpoint_interval_chunks", 1)))
    max_consecutive_errors = int(gen_config.get("max_consecutive_llm_errors", 3))
    consecutive_errors = 0
    for idx, chunk in enumerate(all_chunks, start=1):
        safe_print(f"[{idx}/{len(all_chunks)}] Extracting knowledge: {chunk.chunk_id} {chunk.heading_path}")
        try:
            points = extract_knowledge_points(client, chunk, gen_config)
            all_knowledge_points.extend(points)
            consecutive_errors = 0
        except Exception as exc:  # noqa: BLE001
            consecutive_errors += 1
            add_error(errors, "extract_knowledge", chunk.chunk_id, str(exc))
            safe_print(f"  Failed, skipped: {exc}")
            if not continue_on_error:
                raise
            if max_consecutive_errors > 0 and consecutive_errors >= max_consecutive_errors:
                add_error(
                    errors,
                    "extract_knowledge",
                    "abort_remaining_chunks",
                    f"Stopped after {consecutive_errors} consecutive LLM errors.",
                )
                safe_print(f"  Stopped remaining chunks after {consecutive_errors} consecutive LLM errors.")
                break

        if idx % checkpoint_interval == 0:
            save_progress(
                args.output,
                f"knowledge_chunk_{idx}",
                final_questions,
                candidates,
                all_knowledge_points,
                all_chunks,
                config,
                errors,
            )

    useful_knowledge_points = [
        kp for kp in all_knowledge_points if str(kp.get("is_questionable", "true")).lower() != "false"
    ]
    safe_print(f"Extracted {len(all_knowledge_points)} knowledge points, {len(useful_knowledge_points)} usable.")
    save_progress(
        args.output,
        "knowledge_completed",
        final_questions,
        candidates,
        all_knowledge_points,
        all_chunks,
        config,
        errors,
    )

    chunks_by_doc = group_by(all_chunks, "doc_id")
    knowledge_by_doc = group_by(useful_knowledge_points, "doc_id")
    cross_doc_count = int(gen_config.get("cross_doc_count", 0)) if len(documents) > 1 else 0
    per_doc_budget = max(
        int(gen_config["questions_per_doc_min"]),
        math.ceil(max(1, candidate_budget - cross_doc_count) / max(1, len(documents))),
    )

    for document in documents:
        if len(candidates) >= candidate_budget:
            break
        safe_print(f"Generating candidates for {document.doc_name}")
        remaining_budget = max(0, candidate_budget - len(candidates))
        target_for_doc = min(per_doc_budget, remaining_budget)
        if target_for_doc <= 0:
            break
        try:
            doc_candidates = generate_doc_candidates(
                client,
                document,
                chunks_by_doc.get(document.doc_id, []),
                knowledge_by_doc.get(document.doc_id, []),
                gen_config,
                target_for_doc,
            )
            candidates.extend(doc_candidates[:remaining_budget])
        except Exception as exc:  # noqa: BLE001
            add_error(errors, "generate_doc_candidates", document.doc_id, str(exc))
            safe_print(f"  Failed, skipped: {exc}")
            if not continue_on_error:
                raise

        save_progress(
            args.output,
            f"candidates_{document.doc_id}",
            final_questions,
            candidates,
            all_knowledge_points,
            all_chunks,
            config,
            errors,
        )

    if cross_doc_count > 0 and len(candidates) < candidate_budget:
        safe_print("Generating cross-document candidates")
        try:
            cross_candidates = generate_cross_doc_candidates(client, useful_knowledge_points, gen_config)
            candidates.extend(cross_candidates[: max(0, candidate_budget - len(candidates))])
        except Exception as exc:  # noqa: BLE001
            add_error(errors, "generate_cross_doc_candidates", "cross_doc", str(exc))
            safe_print(f"  Failed, skipped: {exc}")
            if not continue_on_error:
                raise

        save_progress(
            args.output,
            "cross_doc_candidates",
            final_questions,
            candidates,
            all_knowledge_points,
            all_chunks,
            config,
            errors,
        )

    safe_print(f"Generated {len(candidates)} candidates.")

    safe_print("Reviewing and selecting final questions")
    if candidates:
        try:
            final_questions = review_and_select_questions(client, candidates, useful_knowledge_points, gen_config)
        except Exception as exc:  # noqa: BLE001
            add_error(errors, "review_and_select_questions", "final_review", str(exc))
            safe_print(f"  Final review failed, using fallback selection: {exc}")
            if not continue_on_error:
                raise
            final_questions = fallback_select_questions(candidates, target_count)
    else:
        add_error(errors, "review_and_select_questions", "final_review", "No candidates generated.")

    if not final_questions and candidates:
        final_questions = fallback_select_questions(candidates, target_count)

    safe_print(f"Selected {len(final_questions)} final questions.")

    write_excel(args.output, final_questions, candidates, all_knowledge_points, all_chunks, config, errors)
    save_progress(args.output, "completed", final_questions, candidates, all_knowledge_points, all_chunks, config, errors)
    safe_print(f"Wrote Excel: {args.output.resolve()}")
    return 0


def group_by(items: list[Any], key: str) -> dict[str, list[Any]]:
    grouped: dict[str, list[Any]] = {}
    for item in items:
        if isinstance(item, dict):
            value = item.get(key)
        else:
            value = getattr(item, key)
        grouped.setdefault(str(value), []).append(item)
    return grouped


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
